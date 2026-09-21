from __future__ import annotations

import hashlib
import json
import os
import tempfile
import time
from pathlib import Path
from typing import Any

import httpx
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from models import (
    ComplianceOutcome,
    DeliveryAttempt,
    DeliveryRecord,
    FinalResult,
    FinalStatus,
    NotificationStatus,
    utc_now,
)

TRACKER_HEADERS = [
    "Final status",
    "Fit score",
    "Score status",
    "Compliance outcome",
    "Company",
    "Website",
    "Jurisdictions",
    "Employee band",
    "Cloud signals",
    "Fit rationale",
    "Compliance audit",
    "Summary",
    "Contact name",
    "Contact email",
    "Source links",
    "Lead ID",
    "Processed time",
    "Notification status",
]
STATUS_COLORS = {
    FinalStatus.SALES_READY: "C6EFCE",
    FinalStatus.LOWER_PRIORITY: "FFF2CC",
    FinalStatus.REVIEW_REQUIRED: "FCE4D6",
    FinalStatus.DO_NOT_ENGAGE: "FFC7CE",
}


class DeliveryError(RuntimeError):
    pass


def atomic_write_text(path: Path, value: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary = tempfile.mkstemp(dir=path.parent, prefix=f".{path.name}.", suffix=".tmp")
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as handle:
            handle.write(value)
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary, path)
    except BaseException:
        try:
            os.unlink(temporary)
        except FileNotFoundError:
            pass
        raise


def atomic_write_json(path: Path, value: Any) -> None:
    if hasattr(value, "model_dump"):
        value = value.model_dump(mode="json")
    atomic_write_text(path, json.dumps(value, indent=2, ensure_ascii=False) + "\n")


def safe_cell(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    clean = value.replace("\x00", "")
    if clean.startswith(("=", "+", "-", "@")):
        return "'" + clean
    return clean


def _profile_value(result: FinalResult, name: str, default: str = "unknown") -> str:
    if result.accepted_profile is None:
        return default
    claim = getattr(result.accepted_profile, name)
    return str(claim.value)


def _audit_outcome(value: ComplianceOutcome) -> str:
    return value.value.replace("_", " ").title()


def _audit_value(value: Any) -> str:
    return str(value).replace("_", " ").capitalize()


def _audit_percent(value: Any) -> str:
    return f"{float(value):.0%}"


def _named_basis(value: Any) -> str:
    return {
        "likely_alias": "likely the same company",
        "plausible_partial": "possible partial match",
        "distinct_entity": "different company",
        "no_indication": "no evidence of a match",
        "insufficient_identity": "identity not established",
    }.get(str(value), str(value).replace("_", " "))


def _policy_slug(value: str) -> str:
    return "_".join("".join(character.lower() if character.isalnum() else " " for character in value).split())


def _audit_diagnostic(error: str, named_competitors: list[str]) -> str:
    prefix = "missing or malformed Jev answer: competitor_basis__"
    if error.startswith(prefix):
        identifier = error.removeprefix(prefix)
        company = next((name for name in named_competitors if _policy_slug(name) == identifier), identifier.replace("_", " ").title())
        return f"Jev did not return a valid named-competitor decision for {company}."
    if error.startswith("missing or malformed Jev answer: "):
        field = error.removeprefix("missing or malformed Jev answer: ").replace("_", " ")
        return f"Jev did not return a valid answer for {field}."
    return error


def format_compliance_audit(result: FinalResult) -> str:
    sections = [f"Decision: {_audit_outcome(result.compliance_outcome)}"]
    named = [
        check
        for check in result.compliance_checks
        if check.kind == "competitor" and check.policy_entry != "competitive_scope"
    ]
    if named:
        lines = ["Named competitor checks"]
        for check in named:
            line = f"• {check.policy_entry}: {_audit_outcome(check.outcome)}"
            basis = check.details.get("basis")
            if basis:
                line += f", {_named_basis(basis)}"
            else:
                line += f"\n  {check.reason}"
            lines.append(line)
        sections.append("\n".join(lines))

    scope_checks = [check for check in result.compliance_checks if check.policy_entry == "competitive_scope"]
    if scope_checks:
        check = scope_checks[0]
        lines = ["Capability assessment", f"• {_audit_outcome(check.outcome)}"]
        required_details = {
            "openai_classification",
            "jev_score",
            "jev_confidence",
            "material_overlap_probability",
            "score_threshold",
            "decision_probability_threshold",
        }
        if required_details <= check.details.keys():
            lines.extend(
                [
                    f"• OpenAI classification: {_audit_value(check.details['openai_classification'])}",
                    "• Jev overlap: "
                    f"{float(check.details['jev_score']):.2f} / 3.00 "
                    f"({_audit_percent(check.details['jev_confidence'])} confidence)",
                    "• Material-overlap probability: "
                    f"{_audit_percent(check.details['material_overlap_probability'])}",
                    "• Flag rule: score ≥ "
                    f"{float(check.details['score_threshold']):.2f} and probability ≥ "
                    f"{_audit_percent(check.details['decision_probability_threshold'])}",
                ]
            )
        else:
            lines.append(f"  {check.reason}")
        sections.append("\n".join(lines))

    country_checks = [check for check in result.compliance_checks if check.kind == "country"]
    if country_checks:
        lines = ["Country policy"]
        for check in country_checks:
            lines.extend(
                [
                    f"• {_audit_outcome(check.outcome)}",
                    f"  {check.reason}",
                ]
            )
        sections.append("\n".join(lines))

    diagnostic_checks = [check for check in result.compliance_checks if check.kind == "assessment"]
    if diagnostic_checks:
        lines = ["Data-quality notes"]
        competitor_names = [check.policy_entry for check in named]
        for check in diagnostic_checks:
            errors = check.details.get("errors") or [check.reason]
            lines.extend(f"• {_audit_diagnostic(error, competitor_names)}" for error in errors)
        sections.append("\n".join(lines))
    return "\n\n".join(sections)


def tracker_row(result: FinalResult) -> list[Any]:
    cloud_signals = ""
    jurisdictions = ""
    if result.accepted_profile:
        cloud_signals = " | ".join(item.value for item in result.accepted_profile.cloud_signals)
        jurisdictions = " | ".join(
            f"{item.entity_name}: {item.country} ({item.role.value})"
            for item in result.accepted_profile.jurisdictions
        )
    compliance_reason = format_compliance_audit(result)
    company = result.submission.company_name
    if result.synthetic:
        company = f"[SYNTHETIC TEST] {company}"
    return [
        result.final_status.value,
        result.fit.score,
        result.fit.score_status,
        result.compliance_outcome.value,
        company,
        result.submission.website,
        jurisdictions,
        _profile_value(result, "employee_count_band"),
        cloud_signals,
        result.fit.rationale,
        compliance_reason,
        result.summary,
        result.submission.name,
        result.submission.email,
        "\n".join(result.evidence_references),
        result.lead_id,
        result.processed_at.isoformat(),
        result.notification_status.value,
    ]


def update_workbook(path: Path, result: FinalResult) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        workbook = load_workbook(path)
        sheet = workbook.active
        headers = [cell.value for cell in sheet[1]]
        migrated_headers = [
            "Jurisdictions"
            if value == "Headquarters"
            else "Compliance audit"
            if value == "Compliance reason"
            else value
            for value in headers
        ]
        if migrated_headers == TRACKER_HEADERS:
            for index, value in enumerate(TRACKER_HEADERS, start=1):
                sheet.cell(1, index).value = value
            headers = TRACKER_HEADERS
        if headers != TRACKER_HEADERS:
            raise DeliveryError("existing tracker headers do not match the required schema")
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Leads"
        sheet.append(TRACKER_HEADERS)

    lead_column = TRACKER_HEADERS.index("Lead ID") + 1
    target_row = None
    for row_number in range(2, sheet.max_row + 1):
        if sheet.cell(row=row_number, column=lead_column).value == result.lead_id:
            target_row = row_number
            break
    if target_row is None:
        target_row = sheet.max_row + 1
    for column, value in enumerate(tracker_row(result), start=1):
        sheet.cell(row=target_row, column=column).value = safe_cell(value)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(TRACKER_HEADERS))}{sheet.max_row}"
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.cell(target_row, 1).fill = PatternFill("solid", fgColor=STATUS_COLORS[result.final_status])
    widths = [18, 11, 14, 20, 24, 32, 14, 16, 42, 42, 55, 55, 24, 30, 45, 26, 28, 20]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    fd, temporary = tempfile.mkstemp(dir=path.parent, prefix=f".{path.name}.", suffix=".xlsx")
    os.close(fd)
    try:
        workbook.save(temporary)
        os.replace(temporary, path)
    except BaseException:
        try:
            os.unlink(temporary)
        except FileNotFoundError:
            pass
        raise
    finally:
        workbook.close()


def _escape_slack(value: str) -> str:
    value = value.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    return value.replace("@", "＠")


def _display_label(value: Any) -> str:
    return str(value).replace("_", " ").capitalize()


def _primary_compliance_reason(result: FinalResult) -> str:
    if result.compliance_outcome == ComplianceOutcome.CLEAR:
        return "All competitor and country checks cleared under the exercise policy."
    matching = [
        check for check in result.compliance_checks if check.outcome == result.compliance_outcome
    ]
    business_reasons = [check for check in matching if check.kind != "assessment"]
    selected = (business_reasons or matching or result.compliance_checks)[0] if (business_reasons or matching or result.compliance_checks) else None
    if selected is None:
        return f"Compliance is {_display_label(result.compliance_outcome).casefold()}."
    if (
        selected.kind == "competitor"
        and selected.policy_entry == "competitive_scope"
        and selected.outcome == ComplianceOutcome.FLAGGED
    ):
        return "This company offers cloud cost management services that directly compete with us."
    return " ".join(selected.reason.split())


def build_notification(result: FinalResult) -> str:
    score = "Unknown" if result.fit.score is None else str(result.fit.score)
    prefix = "[SYNTHETIC TEST] " if result.synthetic else ""
    status = _display_label(result.final_status)
    score_status = _display_label(result.fit.score_status)
    compliance = _display_label(result.compliance_outcome)
    company = _escape_slack(result.submission.company_name)
    summary = _escape_slack(result.summary)
    reason = _escape_slack(_primary_compliance_reason(result))
    message = (
        f"{prefix}Lead assessment\n"
        f"\nStatus: {status}\n"
        f"Company: {company}\n"
        f"Fit: {score} ({score_status})\n"
        f"\nSummary\n{summary}\n"
        f"\nCompliance\n{compliance}\n{reason}"
    )
    return message


def build_slack_blocks(result: FinalResult) -> list[dict[str, Any]]:
    score = "Unknown" if result.fit.score is None else str(result.fit.score)
    title = "Lead assessment"
    if result.synthetic:
        title = "[SYNTHETIC TEST] Lead assessment"
    reason = _escape_slack(_primary_compliance_reason(result))
    return [
        {
            "type": "header",
            "text": {"type": "plain_text", "text": title},
        },
        {
            "type": "section",
            "fields": [
                {"type": "mrkdwn", "text": f"*Status*\n{_display_label(result.final_status)}"},
                {"type": "mrkdwn", "text": f"*Company*\n{_escape_slack(result.submission.company_name)}"},
                {
                    "type": "mrkdwn",
                    "text": f"*Fit*\n{score} ({_display_label(result.fit.score_status)})",
                },
            ],
        },
        {
            "type": "section",
            "text": {"type": "mrkdwn", "text": f"*Summary*\n{_escape_slack(result.summary)}"},
        },
        {
            "type": "section",
            "text": {
                "type": "mrkdwn",
                "text": f"*Compliance*\n{_display_label(result.compliance_outcome)}\n{reason}",
            },
        },
    ]


def message_hash(message: str) -> str:
    return hashlib.sha256(message.encode("utf-8")).hexdigest()


def load_delivery(path: Path, lead_id: str, message: str) -> DeliveryRecord:
    digest = message_hash(message)
    if not path.exists():
        return DeliveryRecord(lead_id=lead_id, message_sha256=digest)
    try:
        record = DeliveryRecord.model_validate_json(path.read_text(encoding="utf-8"))
    except Exception as exc:
        raise DeliveryError("saved delivery record is invalid") from exc
    if record.lead_id != lead_id or record.message_sha256 != digest:
        return DeliveryRecord(lead_id=lead_id, message_sha256=digest)
    return record


def send_slack(
    webhook_url: str,
    message: str,
    record: DeliveryRecord,
    *,
    explicit_retry: bool = False,
    blocks: list[dict[str, Any]] | None = None,
    client: httpx.Client | None = None,
) -> DeliveryRecord:
    if record.status == NotificationStatus.SENT and record.message_sha256 == message_hash(message):
        return record
    if record.status in {NotificationStatus.FAILED, NotificationStatus.UNKNOWN} and not explicit_retry:
        return record
    owns_client = client is None
    http = client or httpx.Client(timeout=httpx.Timeout(connect=10, read=15, write=15, pool=10))
    duplicate_possible = explicit_retry and record.status == NotificationStatus.UNKNOWN
    payload: dict[str, Any] = {"text": message}
    if blocks:
        payload["blocks"] = blocks
    try:
        for attempt_number in range(2):
            try:
                response = http.post(webhook_url, json=payload)
            except (httpx.ConnectError, httpx.ConnectTimeout, httpx.PoolTimeout) as exc:
                if attempt_number == 0:
                    time.sleep(0.5)
                    continue
                attempt = DeliveryAttempt(
                    status=NotificationStatus.FAILED,
                    error=f"{type(exc).__name__}: connection was not established",
                    duplicate_possible=duplicate_possible,
                )
                record.attempts.append(attempt)
                record.status = attempt.status
                return record
            except (
                httpx.ReadTimeout,
                httpx.WriteTimeout,
                httpx.WriteError,
                httpx.ReadError,
                httpx.RemoteProtocolError,
            ) as exc:
                attempt = DeliveryAttempt(
                    status=NotificationStatus.UNKNOWN,
                    error=f"{type(exc).__name__}: Slack acceptance could not be confirmed",
                    duplicate_possible=True,
                )
                record.attempts.append(attempt)
                record.status = attempt.status
                return record
            if response.status_code == 429 and attempt_number == 0:
                try:
                    delay = min(max(float(response.headers.get("Retry-After", "0.5")), 0), 3)
                except ValueError:
                    delay = 0.5
                time.sleep(delay)
                continue
            acknowledgment = response.text[:500]
            if response.status_code == 200 and acknowledgment.strip() == "ok":
                attempt = DeliveryAttempt(
                    status=NotificationStatus.SENT,
                    acknowledgment=acknowledgment,
                    duplicate_possible=duplicate_possible,
                )
                record.attempts.append(attempt)
                record.status = NotificationStatus.SENT
                record.sent_at = utc_now()
                record.acknowledgment = acknowledgment
                record.message_identifier = None
                return record
            if response.status_code >= 500:
                status = NotificationStatus.UNKNOWN
                error = f"Slack returned HTTP {response.status_code}; acceptance is uncertain"
            else:
                status = NotificationStatus.FAILED
                error = f"Slack returned HTTP {response.status_code}: {acknowledgment}"
            attempt = DeliveryAttempt(
                status=status,
                acknowledgment=acknowledgment,
                error=error,
                duplicate_possible=duplicate_possible or status == NotificationStatus.UNKNOWN,
            )
            record.attempts.append(attempt)
            record.status = status
            return record
    finally:
        if owns_client:
            http.close()
    return record

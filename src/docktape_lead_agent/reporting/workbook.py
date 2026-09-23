from __future__ import annotations

import os
import tempfile
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..domain.results import FinalResult
from ..domain.submissions import ComplianceOutcome, FinalStatus
from ..storage.files import DeliveryError

TRACKER_HEADERS = [
    "Final status",
    "Fit score",
    "Score status",
    "Compliance outcome",
    "Company",
    "Website",
    "Jurisdictions",
    "Employee band",
    "Cloud signals",
    "Fit rationale",
    "Compliance audit",
    "Summary",
    "Contact name",
    "Contact email",
    "Source links",
    "Lead ID",
    "Processed time",
    "Notification status",
]
STATUS_COLORS = {
    FinalStatus.SALES_READY: "C6EFCE",
    FinalStatus.LOWER_PRIORITY: "FFF2CC",
    FinalStatus.REVIEW_REQUIRED: "FCE4D6",
    FinalStatus.DO_NOT_ENGAGE: "FFC7CE",
}


def safe_cell(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    clean = value.replace("\x00", "")
    if clean.startswith(("=", "+", "-", "@")):
        return "'" + clean
    return clean


def _profile_value(result: FinalResult, name: str, default: str = "unknown") -> str:
    if result.accepted_profile is None:
        return default
    claim = getattr(result.accepted_profile, name)
    return str(claim.value)


def _audit_outcome(value: ComplianceOutcome) -> str:
    return value.value.replace("_", " ").title()


def _audit_value(value: Any) -> str:
    return str(value).replace("_", " ").capitalize()


def _audit_percent(value: Any) -> str:
    return f"{float(value):.0%}"


def _named_basis(value: Any) -> str:
    return {
        "likely_alias": "likely the same company",
        "plausible_partial": "possible partial match",
        "distinct_entity": "different company",
        "no_indication": "no evidence of a match",
        "insufficient_identity": "identity not established",
    }.get(str(value), str(value).replace("_", " "))


def _policy_slug(value: str) -> str:
    return "_".join(
        "".join(character.lower() if character.isalnum() else " " for character in value).split()
    )


def _audit_diagnostic(error: str, named_competitors: list[str]) -> str:
    prefix = "missing or malformed Jev answer: competitor_basis__"
    if error.startswith(prefix):
        identifier = error.removeprefix(prefix)
        company = next(
            (name for name in named_competitors if _policy_slug(name) == identifier),
            identifier.replace("_", " ").title(),
        )
        return f"Jev did not return a valid named-competitor decision for {company}."
    if error.startswith("missing or malformed Jev answer: "):
        field = error.removeprefix("missing or malformed Jev answer: ").replace("_", " ")
        return f"Jev did not return a valid answer for {field}."
    return error


def format_compliance_audit(result: FinalResult) -> str:
    sections = [f"Decision: {_audit_outcome(result.compliance_outcome)}"]
    named = [
        check
        for check in result.compliance_checks
        if check.kind == "competitor" and check.policy_entry != "competitive_scope"
    ]
    if named:
        lines = ["Named competitor checks"]
        for check in named:
            line = f"• {check.policy_entry}: {_audit_outcome(check.outcome)}"
            basis = check.details.get("basis")
            if basis:
                line += f", {_named_basis(basis)}"
            else:
                line += f"\n  {check.reason}"
            lines.append(line)
        sections.append("\n".join(lines))

    scope_checks = [
        check for check in result.compliance_checks if check.policy_entry == "competitive_scope"
    ]
    if scope_checks:
        check = scope_checks[0]
        lines = ["Capability assessment", f"• {_audit_outcome(check.outcome)}"]
        required_details = {
            "openai_classification",
            "jev_score",
            "jev_confidence",
            "material_overlap_probability",
            "score_threshold",
            "decision_probability_threshold",
        }
        if required_details <= check.details.keys():
            lines.extend(
                [
                    f"• OpenAI classification: {_audit_value(check.details['openai_classification'])}",
                    "• Jev overlap: "
                    f"{float(check.details['jev_score']):.2f} / 3.00 "
                    f"({_audit_percent(check.details['jev_confidence'])} confidence)",
                    "• Material-overlap probability: "
                    f"{_audit_percent(check.details['material_overlap_probability'])}",
                    "• Flag rule: score ≥ "
                    f"{float(check.details['score_threshold']):.2f} and probability ≥ "
                    f"{_audit_percent(check.details['decision_probability_threshold'])}",
                ]
            )
        else:
            lines.append(f"  {check.reason}")
        sections.append("\n".join(lines))

    country_checks = [check for check in result.compliance_checks if check.kind == "country"]
    if country_checks:
        lines = ["Country policy"]
        for check in country_checks:
            lines.extend(
                [
                    f"• {_audit_outcome(check.outcome)}",
                    f"  {check.reason}",
                ]
            )
        sections.append("\n".join(lines))

    diagnostic_checks = [check for check in result.compliance_checks if check.kind == "assessment"]
    if diagnostic_checks:
        lines = ["Data-quality notes"]
        competitor_names = [check.policy_entry for check in named]
        for check in diagnostic_checks:
            errors = check.details.get("errors") or [check.reason]
            lines.extend(f"• {_audit_diagnostic(error, competitor_names)}" for error in errors)
        sections.append("\n".join(lines))
    return "\n\n".join(sections)


def tracker_row(result: FinalResult) -> list[Any]:
    cloud_signals = ""
    jurisdictions = ""
    if result.accepted_profile:
        cloud_signals = " | ".join(item.value for item in result.accepted_profile.cloud_signals)
        jurisdictions = " | ".join(
            f"{item.entity_name}: {item.country} ({item.role.value})"
            for item in result.accepted_profile.jurisdictions
        )
    compliance_reason = format_compliance_audit(result)
    company = result.submission.company_name
    if result.synthetic:
        company = f"[SYNTHETIC TEST] {company}"
    return [
        result.final_status.value,
        result.fit.score,
        result.fit.score_status,
        result.compliance_outcome.value,
        company,
        result.submission.website,
        jurisdictions,
        _profile_value(result, "employee_count_band"),
        cloud_signals,
        result.fit.rationale,
        compliance_reason,
        result.summary,
        result.submission.name,
        result.submission.email,
        "\n".join(result.evidence_references),
        result.lead_id,
        result.processed_at.isoformat(),
        result.notification_status.value,
    ]


def update_workbook(path: Path, result: FinalResult) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        workbook = load_workbook(path)
        sheet = workbook.active
        headers = [cell.value for cell in sheet[1]]
        migrated_headers = [
            "Jurisdictions"
            if value == "Headquarters"
            else "Compliance audit"
            if value == "Compliance reason"
            else value
            for value in headers
        ]
        if migrated_headers == TRACKER_HEADERS:
            for index, value in enumerate(TRACKER_HEADERS, start=1):
                sheet.cell(1, index).value = value
            headers = TRACKER_HEADERS
        if headers != TRACKER_HEADERS:
            raise DeliveryError("existing tracker headers do not match the required schema")
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Leads"
        sheet.append(TRACKER_HEADERS)

    lead_column = TRACKER_HEADERS.index("Lead ID") + 1
    target_row = None
    for row_number in range(2, sheet.max_row + 1):
        if sheet.cell(row=row_number, column=lead_column).value == result.lead_id:
            target_row = row_number
            break
    if target_row is None:
        target_row = sheet.max_row + 1
    for column, value in enumerate(tracker_row(result), start=1):
        sheet.cell(row=target_row, column=column).value = safe_cell(value)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(TRACKER_HEADERS))}{sheet.max_row}"
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.cell(target_row, 1).fill = PatternFill(
        "solid", fgColor=STATUS_COLORS[result.final_status]
    )
    widths = [18, 11, 14, 20, 24, 32, 14, 16, 42, 42, 55, 55, 24, 30, 45, 26, 28, 20]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    fd, temporary = tempfile.mkstemp(dir=path.parent, prefix=f".{path.name}.", suffix=".xlsx")
    os.close(fd)
    try:
        workbook.save(temporary)
        os.replace(temporary, path)
    except BaseException:
        try:
            os.unlink(temporary)
        except FileNotFoundError:
            pass
        raise
    finally:
        workbook.close()

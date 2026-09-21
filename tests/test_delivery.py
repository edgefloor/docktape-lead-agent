from __future__ import annotations

import httpx
from openpyxl import Workbook, load_workbook

from assessment import calculate_fit
from delivery import TRACKER_HEADERS, build_notification, build_slack_blocks, send_slack, tracker_row, update_workbook
from models import (
    ComplianceCheck,
    ComplianceOutcome,
    DeliveryRecord,
    FinalResult,
    FinalStatus,
)


def make_result(lead, profile, assessment) -> FinalResult:
    return FinalResult(
        lead_id=lead.lead_id,
        submission=lead,
        accepted_profile=profile,
        fit=calculate_fit(assessment),
        compliance_checks=[
            ComplianceCheck(kind="country", policy_entry="exercise_country_set", outcome="clear", reason="Clear under exercise rules.")
        ],
        compliance_outcome=ComplianceOutcome.CLEAR,
        final_status=FinalStatus.SALES_READY,
        summary="Supported summary.",
        evidence_references=["https://example.com/about"],
        policy_version="2026-09-21",
    )


def test_workbook_updates_one_row_and_writes_untrusted_text_as_literal(tmp_path, lead, profile, assessment) -> None:
    result = make_result(lead, profile, assessment)
    result.summary = "=HYPERLINK(\"bad\")"
    path = tmp_path / "lead-tracker.xlsx"
    update_workbook(path, result)
    result.fit.score = 60
    update_workbook(path, result)

    workbook = load_workbook(path, data_only=False)
    sheet = workbook.active
    assert sheet.max_row == 2
    assert [cell.value for cell in sheet[1]] == TRACKER_HEADERS
    assert sheet.cell(2, TRACKER_HEADERS.index("Summary") + 1).value.startswith("'=")
    assert sheet.cell(2, TRACKER_HEADERS.index("Fit score") + 1).value == 60
    workbook.close()


def test_workbook_clears_stale_score_and_labels_synthetic_row(tmp_path, lead, profile, assessment) -> None:
    result = make_result(lead, profile, assessment)
    path = tmp_path / "lead-tracker.xlsx"
    update_workbook(path, result)

    result.fit.score = None
    result.fit.score_status = "provisional"
    result.synthetic = True
    update_workbook(path, result)

    workbook = load_workbook(path, data_only=False)
    sheet = workbook.active
    assert sheet.cell(2, TRACKER_HEADERS.index("Fit score") + 1).value is None
    assert sheet.cell(2, TRACKER_HEADERS.index("Company") + 1).value == "[SYNTHETIC TEST] Example Cloud"
    workbook.close()


def test_workbook_migrates_legacy_headquarters_column(tmp_path, lead, profile, assessment) -> None:
    path = tmp_path / "lead-tracker.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    legacy_headers = [
        "Headquarters"
        if value == "Jurisdictions"
        else "Compliance reason"
        if value == "Compliance audit"
        else value
        for value in TRACKER_HEADERS
    ]
    sheet.append(legacy_headers)
    workbook.save(path)
    workbook.close()

    update_workbook(path, make_result(lead, profile, assessment))

    workbook = load_workbook(path, data_only=False)
    sheet = workbook.active
    assert [cell.value for cell in sheet[1]] == TRACKER_HEADERS
    assert "Example Cloud Inc: US (registered_jurisdiction)" in sheet.cell(
        2, TRACKER_HEADERS.index("Jurisdictions") + 1
    ).value
    workbook.close()


def test_notification_escapes_mentions_and_labels_synthetic(lead, profile, assessment) -> None:
    result = make_result(lead, profile, assessment)
    result.synthetic = True
    result.summary = "Contact @channel <now>"
    message = build_notification(result)
    assert message.startswith("[SYNTHETIC TEST]")
    assert "@channel" not in message
    assert "&lt;now&gt;" in message
    assert "Status: Sales ready" in message
    assert "sales_ready" not in message
    assert "\n\nSummary\n" in message
    assert "\n\nCompliance\n" in message
    assert "Sources" not in message
    assert "https://example.com/about" not in message
    blocks = build_slack_blocks(result)
    assert blocks[0]["text"]["text"] == "[SYNTHETIC TEST] Lead assessment"
    assert "*Status*\nSales ready" in blocks[1]["fields"][0]["text"]
    assert len(blocks) == 4
    assert all("Sources" not in str(block) for block in blocks)
    assert all("https://example.com/about" not in str(block) for block in blocks)


def test_notification_shows_one_primary_compliance_reason(lead, profile, assessment) -> None:
    result = make_result(lead, profile, assessment)
    country_reason = "No verified legal, contracting, parent, or operational-headquarters jurisdiction was established."
    diagnostic = "missing or malformed Jev answer: competitor_basis__spendwise_cloud"
    result.compliance_outcome = ComplianceOutcome.REVIEW_REQUIRED
    result.final_status = FinalStatus.REVIEW_REQUIRED
    result.compliance_checks = [
        ComplianceCheck(
            kind="competitor",
            policy_entry="competitive_scope",
            outcome=ComplianceOutcome.CLEAR,
            reason="Both reviewers found no material competitive overlap.",
        ),
        ComplianceCheck(
            kind="country",
            policy_entry="exercise_country_set",
            outcome=ComplianceOutcome.REVIEW_REQUIRED,
            reason=country_reason,
        ),
        ComplianceCheck(
            kind="assessment",
            policy_entry="jev_response_validation",
            outcome=ComplianceOutcome.REVIEW_REQUIRED,
            reason=diagnostic,
        ),
    ]

    message = build_notification(result)
    blocks = build_slack_blocks(result)

    assert country_reason in message
    assert diagnostic not in message
    assert " | " not in message
    assert country_reason in blocks[-1]["text"]["text"]
    assert diagnostic not in blocks[-1]["text"]["text"]


def test_notification_translates_competitor_flag_for_sales(lead, profile, assessment) -> None:
    result = make_result(lead, profile, assessment)
    technical_reason = (
        "OpenAI identified an unlisted competing service in Cloud cost optimization and FinOps; "
        "Jev independently rated capability overlap 3.00/3.00 at confidence 1.00 "
        "(material-overlap probability 1.00)."
    )
    sales_reason = "This company offers cloud cost management services that directly compete with us."
    result.compliance_outcome = ComplianceOutcome.FLAGGED
    result.final_status = FinalStatus.DO_NOT_ENGAGE
    result.compliance_checks = [
        ComplianceCheck(
            kind="competitor",
            policy_entry="competitive_scope",
            outcome=ComplianceOutcome.FLAGGED,
            reason=technical_reason,
        )
    ]

    message = build_notification(result)
    blocks = build_slack_blocks(result)

    assert sales_reason in message
    assert sales_reason in blocks[-1]["text"]["text"]
    assert "OpenAI" not in message
    assert "Jev" not in message
    assert "confidence" not in message
    assert "probability" not in message


def test_tracker_formats_compliance_as_a_structured_audit(lead, profile, assessment) -> None:
    result = make_result(lead, profile, assessment)
    result.compliance_outcome = ComplianceOutcome.FLAGGED
    result.final_status = FinalStatus.DO_NOT_ENGAGE
    result.compliance_checks = [
        ComplianceCheck(
            kind="competitor",
            policy_entry="CloudTrim Inc",
            outcome=ComplianceOutcome.CLEAR,
            reason="No supported competitor match to CloudTrim Inc. Basis: distinct_entity.",
            details={"basis": "distinct_entity"},
        ),
        ComplianceCheck(
            kind="competitor",
            policy_entry="competitive_scope",
            outcome=ComplianceOutcome.FLAGGED,
            reason=(
                "OpenAI identified an unlisted competing service; Jev independently rated capability "
                "overlap 3.00/3.00 at confidence 1.00 (material-overlap probability 1.00)."
            ),
            details={
                "openai_classification": "competing_service",
                "jev_score": 3.0,
                "jev_confidence": 1.0,
                "material_overlap_probability": 1.0,
                "non_material_overlap_probability": 0.0,
                "score_threshold": 2.0,
                "decision_probability_threshold": 0.8,
            },
        ),
        ComplianceCheck(
            kind="country",
            policy_entry="exercise_country_set",
            outcome=ComplianceOutcome.REVIEW_REQUIRED,
            reason="No verified legal jurisdiction was established.",
        ),
    ]

    audit = tracker_row(result)[TRACKER_HEADERS.index("Compliance audit")]

    assert audit.startswith("Decision: Flagged")
    assert "\n\nNamed competitor checks\n" in audit
    assert "\n\nCapability assessment\n" in audit
    assert "\n\nCountry policy\n" in audit
    assert "CloudTrim Inc: Clear, different company" in audit
    assert "OpenAI classification: Competing service" in audit
    assert "Jev overlap: 3.00 / 3.00 (100% confidence)" in audit
    assert "Material-overlap probability: 100%" in audit
    assert "Flag rule: score ≥ 2.00 and probability ≥ 80%" in audit
    assert "OpenAI identified" not in audit
    assert "FINAL OUTCOME" not in audit
    assert "Evidence:" not in audit
    assert " | " not in audit


def test_slack_acknowledgment_and_unknown_server_failure() -> None:
    def ok(request: httpx.Request) -> httpx.Response:
        return httpx.Response(200, text="ok", request=request)

    record = DeliveryRecord(lead_id="lead", message_sha256="hash")
    with httpx.Client(transport=httpx.MockTransport(ok)) as client:
        sent = send_slack("https://hooks.slack.test/x", "message", record, client=client)
    assert sent.status == "sent"
    assert sent.message_identifier is None

    def failed(request: httpx.Request) -> httpx.Response:
        return httpx.Response(500, text="error", request=request)

    record = DeliveryRecord(lead_id="lead", message_sha256="other")
    with httpx.Client(transport=httpx.MockTransport(failed)) as client:
        uncertain = send_slack("https://hooks.slack.test/x", "changed", record, client=client)
    assert uncertain.status == "unknown"


def test_slack_write_timeout_is_unknown_and_requires_explicit_retry() -> None:
    calls = 0

    def timeout(request: httpx.Request) -> httpx.Response:
        nonlocal calls
        calls += 1
        raise httpx.WriteTimeout("ambiguous", request=request)

    record = DeliveryRecord(lead_id="lead", message_sha256="hash")
    with httpx.Client(transport=httpx.MockTransport(timeout)) as client:
        uncertain = send_slack("https://hooks.slack.test/x", "message", record, client=client)
        unchanged = send_slack("https://hooks.slack.test/x", "message", uncertain, client=client)
    assert uncertain.status == "unknown"
    assert unchanged.status == "unknown"
    assert calls == 1
